"""Convert an uploaded CSV/TSV/XLSX into Parquet using bounded memory."""

import codecs
import csv
import io
import re
import unicodedata
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.csv as pacsv
import pyarrow.parquet as pq
from charset_normalizer import from_bytes
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

ALL_STRINGS_WARNING = (
    "Column types could not be inferred consistently, so every column was stored as text."
)
SAMPLE_BYTES = 1024 * 1024
CSV_BLOCK_BYTES = 1024 * 1024
XLSX_BATCH_ROWS = 10_000
ROW_GROUP_ROWS = 100_000
MAX_COLUMN_NAME = 63
_CANDIDATE_DELIMITERS = (",", ";", "\t", "|")
_CAMEL = re.compile(r"(?<=[a-z0-9])(?=[A-Z])")
_NON_WORD = re.compile(r"[^\w]+", re.UNICODE)


class IngestError(Exception):
    """A user-facing reason why a file could not be converted."""


@dataclass
class IngestResult:
    """Outcome of converting one file."""

    row_count: int
    col_count: int
    column_map: list[dict[str, str]]
    encoding: str | None = None
    delimiter: str | None = None
    sheet: str | None = None
    warnings: list[str] = field(default_factory=list)


def normalize_columns(names: list[str]) -> list[str]:
    """Return unique snake_case names (unicode letters kept) for the given headers."""
    used: set[str] = set()
    result: list[str] = []
    for index, raw in enumerate(names, start=1):
        text = unicodedata.normalize("NFKC", str(raw).strip())
        text = _CAMEL.sub("_", text).lower()
        text = _NON_WORD.sub("_", text).strip("_")[:MAX_COLUMN_NAME].strip("_")
        if not text:
            text = f"col_{index}"
        elif text[0].isdigit():
            text = f"col_{text}"[:MAX_COLUMN_NAME]
        candidate, suffix = text, 2
        while candidate in used:
            tail = f"_{suffix}"
            candidate = f"{text[: MAX_COLUMN_NAME - len(tail)]}{tail}"
            suffix += 1
        used.add(candidate)
        result.append(candidate)
    return result


def _column_map(originals: list[str], normalized: list[str]) -> list[dict[str, str]]:
    return [{"original": o, "normalized": n} for o, n in zip(originals, normalized, strict=True)]


def detect_encoding(sample: bytes) -> str:
    """Guess the text encoding of a sample (BOMs, strict UTF-8, then charset-normalizer)."""
    if sample.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    try:
        codecs.getincrementaldecoder("utf-8")().decode(sample, final=False)
        return "utf-8"
    except UnicodeDecodeError:
        pass
    best = from_bytes(sample).best()
    if best is None:
        raise IngestError("The text encoding of this file could not be determined.")
    return str(best.encoding)


def detect_delimiter(text: str) -> str:
    """Pick the delimiter that splits the first lines into the most consistent columns."""
    lines = [ln for ln in text.splitlines() if ln.strip()][:25]
    best, best_score = ",", (0, 0)
    for candidate in _CANDIDATE_DELIMITERS:
        try:
            counts = [len(row) for row in csv.reader(lines, delimiter=candidate)]
        except csv.Error:
            continue
        if not counts:
            continue
        width = max(set(counts), key=counts.count)
        score = (counts.count(width) if width > 1 else 0, width)
        if score > best_score:
            best, best_score = candidate, score
    return best


def _tmp(path: Path) -> Path:
    return path.with_name(path.name + ".tmp")


def _write_batches(
    batches: Iterable[pa.RecordBatch], dst: Path, names: list[str], max_rows: int
) -> int:
    """Stream record batches into a Parquet file; return the row count."""
    rows = 0
    writer: pq.ParquetWriter | None = None
    tmp = _tmp(dst)
    try:
        for batch in batches:
            if batch.num_rows == 0:
                continue
            rows += batch.num_rows
            if rows > max_rows:
                raise IngestError(f"The file has more than {max_rows:,} rows (the allowed limit).")
            table = pa.Table.from_batches([batch]).rename_columns(names)
            if writer is None:
                writer = pq.ParquetWriter(tmp, table.schema, compression="zstd")
            writer.write_table(table, row_group_size=ROW_GROUP_ROWS)
    except BaseException:
        if writer is not None:
            writer.close()
            writer = None
        tmp.unlink(missing_ok=True)
        raise
    finally:
        if writer is not None:
            writer.close()
    if rows == 0:
        tmp.unlink(missing_ok=True)
        raise IngestError("The file has a header but no data rows.")
    tmp.replace(dst)
    return rows


def _iter_batches(reader: pacsv.CSVStreamingReader) -> Iterator[pa.RecordBatch]:
    while True:
        try:
            yield reader.read_next_batch()
        except StopIteration:
            return


def _open_csv(
    src: Path, encoding: str, delimiter: str, strings_only: tuple[int, int] | None
) -> pacsv.CSVStreamingReader:
    """Open a streaming reader; `strings_only=(skip_rows, ncols)` forces all-text columns."""
    arrow_encoding = "utf-8" if encoding == "utf-8-sig" else encoding
    if strings_only is None:
        read_opts = pacsv.ReadOptions(encoding=arrow_encoding, block_size=CSV_BLOCK_BYTES)
        convert_opts = pacsv.ConvertOptions(strings_can_be_null=True)
    else:
        skip, ncols = strings_only
        cols = [f"c{i}" for i in range(ncols)]
        read_opts = pacsv.ReadOptions(
            encoding=arrow_encoding,
            block_size=CSV_BLOCK_BYTES,
            skip_rows=skip,
            column_names=cols,
        )
        convert_opts = pacsv.ConvertOptions(
            strings_can_be_null=True, column_types={c: pa.string() for c in cols}
        )
    return pacsv.open_csv(
        src,
        read_options=read_opts,
        parse_options=pacsv.ParseOptions(delimiter=delimiter),
        convert_options=convert_opts,
    )


def _ingest_csv(src: Path, dst: Path, ext: str, max_rows: int) -> IngestResult:
    with src.open("rb") as handle:
        sample = handle.read(SAMPLE_BYTES)
    if not sample.strip():
        raise IngestError("The file is empty.")
    encoding = detect_encoding(sample)
    text = sample.decode("utf-16" if encoding == "utf-16" else encoding, errors="replace")
    text = text.lstrip("\ufeff")
    lines = text.splitlines()
    leading_blank = next((i for i, ln in enumerate(lines) if ln != ""), len(lines))
    first = lines[leading_blank] if leading_blank < len(lines) else ""
    if not first.strip(" ,;\t|\"'"):
        raise IngestError("The file has no header row (the first line is empty).")
    delimiter = "\t" if ext == "tsv" else detect_delimiter(text)

    warnings: list[str] = []
    try:
        reader = _open_csv(src, encoding, delimiter, None)
        originals = [n.lstrip("\ufeff") for n in reader.schema.names]
        names = normalize_columns(originals)
        rows = _write_batches(_iter_batches(reader), dst, names, max_rows)
    except pa.ArrowInvalid:
        # Typically a later value does not fit the types inferred from the first block.
        _tmp(dst).unlink(missing_ok=True)
        header = next(
            csv.reader(io.StringIO("\n".join(lines[leading_blank:])), delimiter=delimiter)
        )
        originals = [h.lstrip("\ufeff") for h in header]
        names = normalize_columns(originals)
        warnings.append(ALL_STRINGS_WARNING)
        try:
            reader = _open_csv(src, encoding, delimiter, (leading_blank + 1, len(originals)))
            rows = _write_batches(_iter_batches(reader), dst, names, max_rows)
        except pa.ArrowInvalid as exc:
            _tmp(dst).unlink(missing_ok=True)
            raise IngestError(
                f"The file could not be parsed (inconsistent number of columns?): {exc}"
            ) from exc
    except UnicodeDecodeError as exc:
        _tmp(dst).unlink(missing_ok=True)
        raise IngestError(f"The file is not valid {encoding} text.") from exc
    except IngestError:
        _tmp(dst).unlink(missing_ok=True)
        raise
    return IngestResult(
        row_count=rows,
        col_count=len(originals),
        column_map=_column_map(originals, names),
        encoding=encoding,
        delimiter=delimiter,
        warnings=warnings,
    )


def _open_workbook(src: Path) -> Workbook:
    try:
        return load_workbook(src, read_only=True, data_only=True)
    except Exception as exc:
        raise IngestError("The Excel file could not be opened (it may be corrupted).") from exc


def list_sheets(src: Path) -> list[str]:
    """Return the sheet names of an .xlsx file."""
    workbook = _open_workbook(src)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def _sheet_rows(
    workbook: Workbook, sheet: str | None
) -> tuple[str, list[str], Iterator[tuple[Any, ...]]]:
    """Return (sheet name, header names, iterator over the remaining rows)."""
    if sheet is None:
        sheet = workbook.sheetnames[0] if workbook.sheetnames else None
    if sheet is None or sheet not in workbook.sheetnames:
        raise IngestError(f"Sheet '{sheet}' was not found in the workbook.")
    worksheet = workbook[sheet]
    worksheet.reset_dimensions()
    rows = iter(worksheet.iter_rows(values_only=True))
    for row in rows:
        if any(cell is not None and str(cell).strip() for cell in row):
            names = [str(c).strip() if c is not None else "" for c in row]
            while names and not names[-1]:
                names.pop()
            return sheet, names, rows
    raise IngestError("The selected sheet is empty or has no header row.")


def _xlsx_batches(
    rows: Iterator[tuple[Any, ...]], ncols: int, *, strings_only: bool
) -> Iterator[pa.RecordBatch]:
    """Yield batches; types come from the first batch and later mismatches raise Arrow errors."""
    chunk: list[list[Any]] = []
    schema: pa.Schema | None = None
    labels = [f"c{i}" for i in range(ncols)]

    def build() -> pa.RecordBatch:
        columns = [[r[i] for r in chunk] for i in range(ncols)]
        if strings_only:
            arrays = [pa.array([_cell_text(v) for v in col], type=pa.string()) for col in columns]
        elif schema is None:
            arrays = [pa.array(col) for col in columns]
        else:
            arrays = [pa.array(col, type=schema.field(i).type) for i, col in enumerate(columns)]
        return pa.RecordBatch.from_arrays(arrays, names=labels)

    for row in rows:
        values = list(row[:ncols]) + [None] * (ncols - len(row))
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        chunk.append(values)
        if len(chunk) >= XLSX_BATCH_ROWS:
            batch = build()
            schema = schema or batch.schema
            yield batch
            chunk = []
    if chunk:
        yield build()


def _ingest_xlsx(src: Path, dst: Path, sheet: str | None, max_rows: int) -> IngestResult:
    warnings: list[str] = []
    for strings_only in (False, True):
        workbook = _open_workbook(src)
        try:
            sheet_name, originals, rows = _sheet_rows(workbook, sheet)
            names = normalize_columns(originals)
            batches = _xlsx_batches(rows, len(originals), strings_only=strings_only)
            count = _write_batches(batches, dst, names, max_rows)
        except (pa.ArrowInvalid, pa.ArrowTypeError) as exc:
            _tmp(dst).unlink(missing_ok=True)
            if strings_only:
                raise IngestError(f"The sheet could not be converted: {exc}") from exc
            warnings.append(ALL_STRINGS_WARNING)
            continue
        except IngestError:
            _tmp(dst).unlink(missing_ok=True)
            raise
        finally:
            workbook.close()
        return IngestResult(
            row_count=count,
            col_count=len(originals),
            column_map=_column_map(originals, names),
            sheet=sheet_name,
            warnings=warnings,
        )
    raise IngestError("The sheet could not be converted.")  # pragma: no cover


def ingest_file(
    src: Path, dst: Path, ext: str, *, sheet: str | None = None, max_rows: int
) -> IngestResult:
    """Convert `src` to Parquet at `dst` and return what was found.

    Memory use is bounded by the batch size, not the file size. A type clash in a
    later batch restarts the file with every column as text instead of failing.
    """
    if ext == "xlsx":
        return _ingest_xlsx(src, dst, sheet, max_rows)
    if ext in {"csv", "tsv"}:
        return _ingest_csv(src, dst, ext, max_rows)
    raise IngestError(f"Unsupported file type: {ext}")
